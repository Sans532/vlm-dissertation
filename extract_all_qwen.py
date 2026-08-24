import re
from docx import Document

def extract_all():
    with open("DISSERTATION_RESULTS.md", "r") as f:
        content = f.read()

    start_idx = content.find("### 1.2 Qwen2.5-VL-7B")
    end_idx = content.find("### 1.3 VideoLLaVA")
    
    qwen_section = content[start_idx:end_idx]
    
    results = []

    # 1. Parse tables
    lines = qwen_section.split('\n')
    in_table = False
    headers = []
    for line in lines:
        line = line.strip()
        if line.startswith('|') and not line.startswith('|-'):
            cells = [c.strip() for c in line.split('|')][1:-1]
            if not in_table:
                headers = cells
                in_table = True
            else:
                row_data = dict(zip(headers, cells))
                if 'Overall' in row_data:
                    prompt = row_data.get('Prompt', 'Binary')
                    nframes = row_data.get('Frames', '8')
                    # infer prompt if missing from table context
                    if 'Novice acc' in row_data and 'Other classes' not in row_data and 'Expert acc' in row_data:
                        prompt = 'Binary'
                    elif 'Other classes' in row_data and prompt == 'Binary':
                        if 'reasoning' in row_data.get('File', ''): prompt = 'Reasoning'
                        elif 'fourclass' in row_data.get('File', ''): prompt = 'Fourclass'
                        elif 'structured' in row_data.get('File', ''): prompt = 'Structured'

                    trim = row_data.get('Trim', 'entire')
                    if 'trimmed' in row_data.get('File', ''): trim = 'trimmed'
                    view = row_data.get('View', 'ego') # fallback
                    if 'exo' in row_data.get('File', '') and view == 'ego': view = 'exo'
                    
                    overall_str = row_data.get('Overall', '')
                    acc_match = re.search(r'([\d\.]+%)', overall_str)
                    accuracy = acc_match.group(1) if acc_match else overall_str

                    novice = row_data.get('Novice acc', '')
                    other = row_data.get('Other classes', '')
                    expert = row_data.get('Expert acc', '')
                    recall = f"Novice: {novice}"
                    if expert: recall += f", Expert: {expert}"
                    if other: recall += f", Other: {other}"
                    
                    dist = row_data.get('Predicted counts', '')
                    filename = row_data.get('File', '').replace('`', '')
                    
                    results.append({
                        'Prompts': prompt + ' (Collapsed)',
                        'nframes': nframes,
                        'Trim': trim,
                        'view': view,
                        'accuracy': accuracy,
                        'precision': '',
                        'recall': recall,
                        'label distribution': dist,
                        'file': filename
                    })
        elif line.startswith('|-'):
            pass
        else:
            in_table = False

    # 2. Parse the bullet lists ("Kept")
    blocks = re.split(r'\n\*\*(.*?)\*\*\s*\n', qwen_section)
    for i in range(1, len(blocks), 2):
        header = blocks[i].strip()
        body = blocks[i+1]
        
        if '—' not in header: continue
        parts = [p.strip() for p in header.replace('—', ',').split(',')]
        prompt = parts[0]
        nframes = "native"
        for p in parts:
            if 'frames' in p: nframes = p.replace('frames', '').strip()
            elif 'best-exo' in p: nframes = '8'
        trim = "trimmed" if "trimmed" in header else "entire"
        view = "ego" if "ego" in header else "exo"
        if "best-exo" in header: view = "best-exo"
        
        if "n=389" in header: prompt += " (n=389)"

        subruns = []
        if re.search(r'- ego:\s*Overall', body) and re.search(r'- exo:\s*Overall', body):
            # Parse shared file at the end
            shared_file = ""
            for lb in body.strip().split('\n'):
                if lb.strip().startswith('- File:'):
                    shared_file = lb.replace('- File:', '').strip().strip('`')

            for subview in ['ego', 'exo']:
                sub_match = re.search(rf'- {subview}:.*?Overall.*?= \*\*([\d\.]+%)\*\* — (.*?)\. Predicted counts: (.*)', body)
                if sub_match:
                    subruns.append({
                        'view': subview,
                        'accuracy': sub_match.group(1),
                        'recall': sub_match.group(2).strip(),
                        'dist': sub_match.group(3).strip(),
                        'file': shared_file
                    })
        elif '- Overall:' in body:
            acc_m = re.search(r'- Overall:.*?=\s*\*\*?([\d\.]+%)', body)
            acc = acc_m.group(1) if acc_m else ""
            
            lines_body = body.strip().split('\n')
            recall = ""
            dist = ""
            filename = ""
            for lb in lines_body:
                lb = lb.strip()
                if lb.startswith('- Novice:') or lb.startswith('- Intermediate Expert:') or lb.startswith('- Early Expert:'):
                    recall = lb.lstrip('- ').strip()
                elif lb.startswith('- Predicted counts:'):
                    dist = lb.replace('- Predicted counts:', '').strip()
                elif lb.startswith('- File:'):
                    filename = lb.replace('- File:', '').strip().strip('`')
            
            subruns.append({
                'view': view,
                'accuracy': acc,
                'recall': recall,
                'dist': dist,
                'file': filename
            })
            
        for subrun in subruns:
            results.append({
                'Prompts': prompt,
                'nframes': nframes,
                'Trim': trim,
                'view': subrun['view'],
                'accuracy': subrun['accuracy'],
                'precision': '',
                'recall': subrun['recall'],
                'label distribution': subrun['dist'],
                'file': subrun['file']
            })

    return results

def main():
    results = extract_all()
    # Create new docx
    doc = Document()
    doc.add_heading('Qwen2.5-VL-7B Climbing Results', 0)
    table = doc.add_table(rows=1, cols=9)
    table.style = 'Table Grid'
    headers = ['Prompts', 'nframes', 'Trim', 'view', 'accuracy', 'precision', 'recall', 'label distribution', 'File']
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
        
    for res in results:
        cells = table.add_row().cells
        cells[0].text = res['Prompts']
        cells[1].text = res['nframes']
        cells[2].text = res['Trim']
        cells[3].text = res['view']
        cells[4].text = res['accuracy']
        cells[5].text = res['precision']
        cells[6].text = res['recall']
        cells[7].text = res['label distribution']
        cells[8].text = res['file']
        
    doc.save('Qwen2.5-VL-7B_Climbing_results.docx')
    print(f"Total rows written: {len(results)}")

if __name__ == '__main__':
    main()
